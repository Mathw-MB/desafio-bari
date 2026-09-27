"""
relatorio_excel.py
Desafio Prático - Estágio AI & Data Lab | Bari - Parte 2 (Automação/RPA)

Exportador Excel. Consome o mesmo dicionário de metricas.calcular_metricas()
que relatorio_html.py - opção secundária de formato para quem prefere abrir
os números direto numa planilha (filtrar, montar tabela dinâmica, etc.) em
vez de um relatório fechado. Depende só de pandas + openpyxl.
"""

from datetime import datetime

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COR_CABECALHO = "1C5CAB"   # azul sequencial escuro da paleta
COR_AVISO = "FFF3CD"
COR_CRITICO = "F8D7DA"


def _formatar_cabecalho(ws, n_colunas):
    for col in range(1, n_colunas + 1):
        celula = ws.cell(row=1, column=col)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor=COR_CABECALHO)
        celula.alignment = Alignment(horizontal="left")


def _autoajustar_colunas(ws, df):
    for i, coluna in enumerate(df.columns, start=1):
        largura = max(len(str(coluna)), df[coluna].astype(str).map(len).max() if len(df) else 0) + 3
        ws.column_dimensions[get_column_letter(i)].width = min(largura, 45)


def gerar_excel(metricas: dict, caminho_saida: str, nome_arquivo_origem: str) -> None:
    conv_geral = metricas["conversao_geral"]
    conv_canal = metricas["conversao_por_canal"].copy()
    funil = metricas["funil_por_etapa"].copy()
    ltv = metricas["ltv"]
    consistencia = metricas["consistencia_etapa_status"]
    conversao_mensal = metricas["conversao_mensal"]
    perda_por_motivo = metricas["perda_por_motivo"].copy()
    qualidade = metricas["qualidade_dados"]

    conv_canal["taxa_conversao"] = conv_canal["taxa_conversao"]
    funil["taxa_queda_condicional"] = funil["taxa_queda_condicional"]

    if conversao_mensal["disponivel"]:
        df_mensal = conversao_mensal["serie"][["mes_rotulo", "total_propostas", "total_contratadas", "taxa_conversao"]].copy()
        df_mensal.columns = ["mes", "total_propostas", "total_contratadas", "taxa_conversao"]
    else:
        df_mensal = pd.DataFrame({"aviso": ["Indisponível nesta execução - coluna 'data_entrada' ausente ou sem data válida."]})

    df_resumo = pd.DataFrame(
        [
            ("Data de geração", datetime.now().strftime("%d/%m/%Y %H:%M")),
            ("Arquivo de origem", nome_arquivo_origem),
            ("Total de propostas", conv_geral["total_propostas"]),
            ("Total contratadas", conv_geral["total_contratadas"]),
            ("Taxa de conversão geral", conv_geral["taxa_conversao"]),
            (f'% propostas com LTV acima de {int(ltv["limite"]*100)}%', ltv["pct_acima_do_limite"]),
            ("Propostas com LTV acima do limite", ltv["propostas_acima_do_limite"]),
            ("etapa_max_funil == 6 consistente com status_final == Contratada",
             "Sim" if consistencia["consistente"] else "Não"),
        ],
        columns=["Métrica", "Valor"],
    )

    df_qualidade = pd.DataFrame(
        [
            ("Encoding do arquivo lido", qualidade["encoding_usado"]),
            ("Linhas brutas", qualidade["linhas_brutas"]),
            ("Linhas tratadas", qualidade["linhas_tratadas"]),
            ("Idades inválidas (nulificadas)", qualidade["linhas_idade_invalida"]),
            ("Datas inconsistentes (assinatura antes da entrada)", qualidade["linhas_data_inconsistente"]),
            ("etapa_max_funil corrigida automaticamente", qualidade["linhas_etapa_corrigida"]),
            ("etapa_max_funil anômala não corrigida (revisão humana)", qualidade["linhas_etapa_anomala_nao_corrigida"]),
            ("Propostas com tipo_imovel = Terreno mantidas", qualidade["propostas_terreno_mantidas"]),
            ("Canais fora das 5 grafias oficiais conhecidas",
             ", ".join(qualidade["canais_nao_reconhecidos"]) or "nenhum"),
        ],
        columns=["Item", "Valor"],
    )
    df_avisos = pd.DataFrame({"Avisos desta execução": qualidade["avisos"] or ["Nenhum aviso nesta execução."]})

    with pd.ExcelWriter(caminho_saida, engine="openpyxl") as writer:
        df_resumo.to_excel(writer, sheet_name="Resumo", index=False)
        df_mensal.to_excel(writer, sheet_name="Conversao_Mensal", index=False)
        conv_canal.to_excel(writer, sheet_name="Conversao_por_Canal", index=False)
        funil.to_excel(writer, sheet_name="Funil_por_Etapa", index=False)
        perda_por_motivo.to_excel(writer, sheet_name="Perda_por_Motivo", index=False)
        df_qualidade.to_excel(writer, sheet_name="Qualidade_de_Dados", index=False)
        df_avisos.to_excel(writer, sheet_name="Qualidade_de_Dados", index=False, startrow=len(df_qualidade) + 3)

        livro = writer.book

        ws_resumo = livro["Resumo"]
        _formatar_cabecalho(ws_resumo, 2)
        _autoajustar_colunas(ws_resumo, df_resumo)
        ws_resumo["B5"].number_format = "0.0%"
        ws_resumo["B6"].number_format = "0.0%"
        if ltv["pct_acima_do_limite"] > 0:
            ws_resumo["B6"].fill = PatternFill("solid", fgColor=COR_CRITICO)
        if not consistencia["consistente"]:
            ws_resumo["B8"].fill = PatternFill("solid", fgColor=COR_CRITICO)

        ws_mensal = livro["Conversao_Mensal"]
        _formatar_cabecalho(ws_mensal, len(df_mensal.columns))
        _autoajustar_colunas(ws_mensal, df_mensal)
        if conversao_mensal["disponivel"] and "taxa_conversao" in df_mensal.columns:
            col_taxa_mensal = df_mensal.columns.get_loc("taxa_conversao") + 1
            for linha in range(2, len(df_mensal) + 2):
                ws_mensal.cell(row=linha, column=col_taxa_mensal).number_format = "0.0%"

        ws_motivo = livro["Perda_por_Motivo"]
        _formatar_cabecalho(ws_motivo, len(perda_por_motivo.columns) or 1)
        if not perda_por_motivo.empty:
            _autoajustar_colunas(ws_motivo, perda_por_motivo)
            col_valor_motivo = perda_por_motivo.columns.get_loc("valor_solicitado_perdido") + 1
            for linha in range(2, len(perda_por_motivo) + 2):
                ws_motivo.cell(row=linha, column=col_valor_motivo).number_format = "R$ #,##0.00"

        ws_canal = livro["Conversao_por_Canal"]
        _formatar_cabecalho(ws_canal, len(conv_canal.columns))
        _autoajustar_colunas(ws_canal, conv_canal)
        col_taxa = conv_canal.columns.get_loc("taxa_conversao") + 1
        for linha in range(2, len(conv_canal) + 2):
            ws_canal.cell(row=linha, column=col_taxa).number_format = "0.0%"

        ws_funil = livro["Funil_por_Etapa"]
        _formatar_cabecalho(ws_funil, len(funil.columns))
        _autoajustar_colunas(ws_funil, funil)
        col_taxa_queda = funil.columns.get_loc("taxa_queda_condicional") + 1
        col_valor = funil.columns.get_loc("valor_solicitado_perdido") + 1
        for linha in range(2, len(funil) + 2):
            ws_funil.cell(row=linha, column=col_taxa_queda).number_format = "0.0%"
            ws_funil.cell(row=linha, column=col_valor).number_format = "R$ #,##0.00"

        ws_qualidade = livro["Qualidade_de_Dados"]
        _formatar_cabecalho(ws_qualidade, 2)
        _autoajustar_colunas(ws_qualidade, df_qualidade)
        if qualidade["avisos"]:
            for linha in range(len(df_qualidade) + 5, len(df_qualidade) + 5 + len(df_avisos)):
                ws_qualidade.cell(row=linha, column=1).fill = PatternFill("solid", fgColor=COR_AVISO)
